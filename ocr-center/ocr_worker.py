# -*- coding: utf-8 -*-
# ocr_worker.py —— PaddleOCR-VL / PP-DocVQA 文档解析（流式版 v3）
# 兼容新旧两种结果结构；每页解析完立即写盘；优先使用官方 Markdown 转换器。
# 用法: python ocr_worker.py <pdf1> [pdf2 ...]
# 保存编码：UTF-8

import os, re, sys, json, time, datetime

BASE     = os.path.dirname(os.path.abspath(__file__))                       # 脚本所在目录
OUT_ROOT = os.environ.get("OCR_OUT_ROOT") or os.path.join(BASE, "output")   # 输出总目录（可用环境变量覆盖）
OUT_DIR  = os.path.join(OUT_ROOT, "ocr")                                    # OCR 输出
IGNORE_LABELS = ["number", "footnote", "vision_footnote", "header", "header_image",
                 "footer", "footer_image", "aside_text"]
CROP_SCALE = 2.0

# 产线候选，按顺序尝试。想优先使用轻量模型，可把 "PaddleOCR-VL" 挪到第一位。
PIPELINE_CANDIDATES = ["PaddleOCR-VL-1.6", "PaddleOCR-VL-1.5", "PaddleOCR-VL",
                       "PP-DocVQA", "PP-DocVQA_v2"]

def log(*a):
    print(*a, flush=True)
    
def json_default(o):
    import numpy as np
    if isinstance(o, np.ndarray):
        return o.tolist()
    if isinstance(o, np.integer):
        return int(o)
    if isinstance(o, np.floating):
        return float(o)
    if isinstance(o, np.bool_):
        return bool(o)
    return str(o)

def create_pipe():
    import paddlex
    log("[版本] paddlex " + str(getattr(paddlex, "__version__", "未知")))
    last_err = None
    for name in PIPELINE_CANDIDATES:
        try:
            pipe = paddlex.create_pipeline(pipeline=name)
            # 3.7.2 默认每批 64 页，攒满才吐结果，看起来像卡死。
            # 关键：外层管线的属性通过 __getattr__ 委托给内部 _pipeline，
            # 必须改到 pipe._pipeline.batch_sampler 上才真正生效。
            try:
                from paddlex.inference.common.batch_sampler import ImageBatchSampler
                inner = getattr(pipe, "_pipeline", pipe)
                inner.batch_sampler = ImageBatchSampler(batch_size=1)
                bs = getattr(inner.batch_sampler, "batch_size", "?")
                log("[OK] 使用产线: %s（内层每批 %s 页，逐页流式输出）" % (name, bs))
            except Exception as e2:
                log("[OK] 使用产线: %s（沿用默认批大小，%s）" % (name, str(e2)[:80]))
            return pipe
        except Exception as e:
            msg = str(e)
            last_err = msg
            log("[尝试] 产线 %s 不可用: %s" % (name, msg[:120]))
            if "does not exist" not in msg and "not supported" not in msg:
                raise
    raise RuntimeError("所有候选产线均不可用，最后错误: " + str(last_err))

def predict_iter(pipe, pdf):
    # 逐页生成器；关闭异步队列避免预渲染多页大图吃光内存；参数逐级降级。
    attempts = [
        dict(use_queues=False, use_layout_detection=True, use_chart_recognition=False,
             use_seal_recognition=False, use_ocr_for_image_block=False,
             format_block_content=False, merge_layout_blocks=True,
             markdown_ignore_labels=list(IGNORE_LABELS)),
        dict(use_queues=False, use_layout_detection=True,
             markdown_ignore_labels=list(IGNORE_LABELS)),
        dict(use_queues=False, use_layout_detection=True),
        dict(use_queues=False),
        {},
    ]
    last_err = None
    for kw in attempts:
        gen = pipe.predict(pdf, **kw)
        try:
            first = next(gen)
        except StopIteration:
            log("[警告] 该参数组合返回 0 页，换参数重试")
            continue
        except TypeError as e:
            last_err = e
            log("[警告] 参数不兼容，降级重试: %s" % str(e)[:100])
            continue
        yield first
        yield from gen
        return
    raise RuntimeError("predict 调用失败: " + str(last_err))

def normalize_page(page_res):
    """新版返回 PaddleOCRVLResult 对象，旧版返回 dict；统一转成可 json 化的 dict。"""
    if isinstance(page_res, dict):
        return page_res
    if hasattr(page_res, "json"):
        try:
            d = page_res.json
            if isinstance(d, dict):
                return d
        except Exception:
            pass
    return page_res

def normalize_block(b, idx):
    """新版 block 是 PaddleOCRVLBlock 对象，旧版是 dict；统一转成 dict。"""
    if isinstance(b, dict):
        d = dict(b)
        d.setdefault("block_id", idx)
        return d
    return {
        "block_label": getattr(b, "label", "") or "",
        "block_content": getattr(b, "content", "") or "",
        "block_bbox": list(getattr(b, "bbox", None) or []),
        "block_id": idx,
    }

def ordered_blocks(page_res):
    """解析结果列表本身就是阅读顺序，直接按序取出并规范化。"""
    bl = page_res.get("parsing_res_list") or []
    return [normalize_block(b, i) for i, b in enumerate(bl)]

def page_width(page_res):
    w = page_res.get("width", 1008)
    return w[0] if isinstance(w, list) else w

def save_md_images(md, imgs_dir):
    """官方转换器返回的 markdown_images: {相对路径: PIL图片}，落到 imgs 目录。"""
    images = md.get("markdown_images") or {}
    for rel, img in images.items():
        try:
            name = os.path.basename(str(rel).replace("\\", "/"))
            if not name:
                continue
            dst = os.path.join(imgs_dir, name)
            if hasattr(img, "save"):
                img.save(dst)
                continue
            import numpy as np
            from PIL import Image
            if isinstance(img, np.ndarray):
                Image.fromarray(img).save(dst)
        except Exception:
            pass

def head_level(t):
    if re.match(r"^\d+\.\d+\.\d+", t):
        return 4
    if re.match(r"^\d+\.\d+", t):
        return 3
    return 2

def save_crop(page_pil, box, dst):
    try:
        if page_pil is None:
            return
        x1, y1, x2, y2 = [int(v * CROP_SCALE) for v in box]
        w, h = page_pil.size
        x1, y1 = max(0, x1), max(0, y1)
        x2, y2 = min(w, x2), min(h, y2)
        if x2 > x1 and y2 > y1:
            page_pil.crop((x1, y1, x2, y2)).save(dst, quality=90)
    except Exception:
        pass

def block_md(b, page_w, imgs_dir, page_pil):
    label = b.get("block_label", "")
    content = (b.get("block_content") or "").strip()
    if not content and label not in ("image", "chart"):
        return ""
    if label in ("display_formula", "formula", "equation"):
        return content if "$" in content else "$$" + content + "$$"
    if label == "inline_formula":
        return content if "$" in content else "$" + content + "$"
    if label == "doc_title":
        return "# " + content
    if label == "paragraph_title":
        return "#" * head_level(content) + " " + content
    if label in ("image", "chart"):
        bb = b.get("block_bbox")
        if not bb or len(bb) < 4:
            return ""
        x1, y1, x2, y2 = [int(v) for v in bb[:4]]
        w_pct = max(1, round((x2 - x1) * 100.0 / page_w))
        fname = "img_in_%s_box_%d_%d_%d_%d.jpg" % (label, x1, y1, x2, y2)
        save_crop(page_pil, (x1, y1, x2, y2), os.path.join(imgs_dir, fname))
        return '<div style="text-align: center;"><img src="imgs/%s" alt="Image" width="%d%%" /></div>' % (fname, w_pct)
    if label in ("figure_title", "image_caption") or re.match(r"^FIGURE", content):
        return '<div style="text-align: center;">%s</div>' % content
    return content

def build_page_md(page_res, blocks, page_w, imgs_dir, pdf, idx):
    """优先用官方转换器；不可用则回退到自建逻辑。"""
    md = getattr(page_res, "markdown", None)
    if isinstance(md, dict) and (md.get("markdown_texts") or md.get("markdown_images")):
        save_md_images(md, imgs_dir)
        return (md.get("markdown_texts") or "").strip()
    # 回退：自建 markdown（旧版 dict 结果）
    labels = [b.get("block_label") for b in blocks]
    has_crops = "image" in labels or "chart" in labels
    page_pil = None
    if has_crops:
        try:
            bmp = pdf[idx].render(scale=CROP_SCALE)
            page_pil = bmp.to_pil().convert("RGB")
        except Exception:
            pass
    parts = []
    for b in blocks:
        if b.get("block_label") in IGNORE_LABELS:
            continue
        s = block_md(b, page_w, imgs_dir, page_pil)
        if s:
            parts.append(s)
    return "\n\n".join(parts)

def add_blocks_docx(doc, blocks, imgs_dir):
    from docx.shared import Inches
    for b in blocks:
        label = b.get("block_label", "")
        if label in IGNORE_LABELS:
            continue
        content = (b.get("block_content") or "").strip()
        if label == "doc_title":
            doc.add_heading(content, level=1)
        elif label == "paragraph_title":
            doc.add_heading(content, level=min(4, head_level(content)))
        elif label in ("display_formula", "inline_formula", "formula", "equation"):
            doc.add_paragraph(content)
        elif label in ("image", "chart"):
            bb = b.get("block_bbox")
            if not bb:
                continue
            fname = "img_in_%s_box_%d_%d_%d_%d.jpg" % (label, int(bb[0]), int(bb[1]), int(bb[2]), int(bb[3]))
            path = os.path.join(imgs_dir, fname)
            if os.path.exists(path):
                p = doc.add_paragraph()
                p.alignment = 1
                p.add_run().add_picture(path, width=Inches(4.0))
        elif content:
            for para in content.split("\n"):
                if para.strip():
                    doc.add_paragraph(para)

def process_one(pipe, pdf_path):
    stem = os.path.splitext(os.path.basename(pdf_path))[0]
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out = os.path.join(OUT_DIR, "%s_%s" % (stem, ts))
    imgs = os.path.join(out, "imgs")
    os.makedirs(imgs, exist_ok=True)
    log("[目录] " + out)
    log("[OCR] 开始解析: " + os.path.basename(pdf_path))

    import pypdfium2 as pdfium
    from docx import Document

    t0 = time.time()
    pdf = pdfium.PdfDocument(pdf_path)
    total = len(pdf)
    page_w = 1008
    docx_all = Document()

    md_f = open(os.path.join(out, "全书合并.md"), "w", encoding="utf-8")
    js_f = open(os.path.join(out, "全书合并.json"), "w", encoding="utf-8")
    js_f.write("[")
    first_page = True
    count = 0
    try:
        for idx, page_res in enumerate(predict_iter(pipe, pdf_path)):
            t1 = time.time()
            if count == 0:
                page_w = page_width(page_res)
            blocks = ordered_blocks(page_res)
            page_md = build_page_md(page_res, blocks, page_w, imgs, pdf, idx)

            # 本页三件套：立刻写盘
            base = "%s_%d" % (stem, idx)
            with open(os.path.join(out, base + ".md"), "w", encoding="utf-8") as f:
                f.write(page_md)
            page_dict = normalize_page(page_res)
            with open(os.path.join(out, base + ".json"), "w", encoding="utf-8") as f:
                json.dump(page_dict, f, ensure_ascii=False, indent=2,default=json_default)
            pd = Document()
            add_blocks_docx(pd, blocks, imgs)
            pd.save(os.path.join(out, base + ".docx"))

            # 合并文件：追加写 + 立刻刷盘
            md_f.write("==== Page %d (source: %s_%d.md) ====\n\n%s\n" % (idx + 1, stem, idx, page_md))
            md_f.flush()
            if not first_page:
                js_f.write(",\n")
            js_f.write(json.dumps(page_dict, ensure_ascii=False, default=json_default))
            js_f.flush()
            first_page = False

            add_blocks_docx(docx_all, blocks, imgs)
            if count > 0:
                docx_all.add_page_break()
            count += 1
            log("[OCR] %s 第 %d/%d 页完成 · 单页 %.1fs · 累计 %.1fs"
                % (stem, idx + 1, total, time.time() - t1, time.time() - t0))
    finally:
        js_f.write("\n]")
        js_f.close()
        md_f.close()
        pdf.close()
        try:
            docx_all.save(os.path.join(out, "全书合并.docx"))
        except Exception as e:
            log("[错误] 合并 docx 保存失败: %s" % e)
    log("[完成] " + out + " · 共 %d 页 · 总用时 %.1f s" % (count, time.time() - t0))

def main():
    if len(sys.argv) < 2:
        log("[错误] 用法: ocr_worker.py <pdf1> [pdf2 ...]")
        sys.exit(1)
    pdfs = [a for a in sys.argv[1:] if os.path.isfile(a)]
    if not pdfs:
        log("[错误] 未找到有效的 PDF 文件")
        sys.exit(1)
    try:
        pipe = create_pipe()
    except Exception as e:
        log("[错误] 创建解析产线失败（请确认该 Python 环境已安装 PaddleX）：", e)
        sys.exit(1)
    log("[提示] 正在加载版面检测 + OCR 大模型，请耐心等待…")
    t_all = time.time()
    for i, pdf in enumerate(pdfs):
        log("[OCR] （%d/%d）%s" % (i + 1, len(pdfs), os.path.basename(pdf)))
        try:
            process_one(pipe, pdf)
        except Exception as e:
            import traceback
            log("[错误] 处理失败:", pdf, e)
            traceback.print_exc()
    log("[完成] 全部 %d 个 PDF 处理完毕，总用时 %.1f s" % (len(pdfs), time.time() - t_all))

if __name__ == "__main__":
    main()
